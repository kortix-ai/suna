#!/usr/bin/env python3
"""Render the synthetic demo dataset: demo-data/src/*.md -> demo-data/dist/ (PDF, DOCX or TXT) + manifest.json.

Run from anywhere: python3 demo-data/build.py. Needs reportlab and python-docx (dev machine only;
the outputs are committed, so seeding never needs them). Every document is fictional.
"""
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent
SRC, DIST = ROOT / "src", ROOT / "dist"
FOOTER = "SYNTHETIC DEMO DOCUMENT. Fictional parties. Not legal advice."

# filename -> (output type, document type)
DOCS = {
    "mutual-nda-quillon": ("docx", "nda"),
    "one-way-nda-harbor": ("pdf", "nda"),
    "msa-brightwater-quillon": ("pdf", "msa"),
    "saas-agreement-lattice": ("docx", "saas"),
    "office-lease-cedar-row": ("pdf", "lease"),
    "employment-agreement-tarrow": ("docx", "employment"),
    "spa-excerpt-aster": ("pdf", "spa"),
    "supply-agreement-meridian": ("pdf", "supply"),
    "dpa-lattice": ("docx", "dpa"),
    "consulting-agreement-pell": ("txt", "consulting"),
}


def blocks(md: str):
    """(kind, text) per block: kind is h1, h2 or p."""
    for raw in re.split(r"\n\s*\n", md.strip()):
        text = " ".join(raw.split())
        if text.startswith("## "):
            yield "h2", text[3:]
        elif text.startswith("# "):
            yield "h1", text[2:]
        else:
            yield "p", text


def esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_pdf(md: str, out: Path) -> int:
    styles = {
        "h1": ParagraphStyle("h1", fontName="Times-Bold", fontSize=16, leading=20, spaceAfter=10),
        "h2": ParagraphStyle("h2", fontName="Times-Bold", fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=4),
        "p": ParagraphStyle("p", fontName="Times-Roman", fontSize=11, leading=15, spaceAfter=6),
    }
    pages = {"n": 0}

    def footer(canvas, doc):
        pages["n"] = doc.page
        canvas.setFont("Helvetica", 7.5)
        canvas.drawCentredString(A4[0] / 2, 28, f"{FOOTER}   Page {doc.page}")

    # Wide margins keep each contract at 2-4 pages, like a real short agreement.
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=110, rightMargin=110, topMargin=90, bottomMargin=80,
                            title=out.stem, author="Lumen demo data", invariant=1)
    story = []
    for kind, text in blocks(md):
        story += [Paragraph(esc(text), styles[kind])]
    story.append(Spacer(1, 6))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return pages["n"]


def render_docx(md: str, out: Path) -> None:
    doc = Document()
    doc.core_properties.author = "Lumen demo data"
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    for kind, text in blocks(md):
        if kind == "h1":
            doc.add_heading(text, level=0)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        else:
            doc.add_paragraph(text)
    doc.sections[0].footer.paragraphs[0].text = FOOTER
    doc.save(str(out))


def main() -> None:
    DIST.mkdir(exist_ok=True)
    (DIST / "templates").mkdir(exist_ok=True)
    manifest = []
    for stem, (kind, doc_type) in DOCS.items():
        md = (SRC / f"{stem}.md").read_text()
        out = DIST / f"{stem}.{kind}"
        pages = None
        if kind == "pdf":
            pages = render_pdf(md, out)
        elif kind == "docx":
            render_docx(md, out)
        else:
            out.write_text("\n\n".join(t for _, t in blocks(md)) + f"\n\n{FOOTER}\n")
        manifest.append({"filename": out.name, "type": doc_type, "pages": pages,
                         "sha256": hashlib.sha256(out.read_bytes()).hexdigest(), "bytes": out.stat().st_size})
    for name in ["playbook-standard-nda.md", "templates/mutual-nda.md", "templates/engagement-letter.md"]:
        (DIST / name).write_text((SRC / name).read_text())
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps([(m["filename"], m["pages"], m["bytes"]) for m in manifest]))


if __name__ == "__main__":
    main()
