"""Tracked-change redlines in DOCX: each edit becomes w:del + w:ins by "Lumen", with a comment.

A PDF or TXT source has no DOCX to mark up, so the redline goes against a
converted copy: a new DOCX written from the document's anchored text.
"""
import copy
import datetime
import os
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .core import KitError, load_doc, normalize

AUTHOR = "Lumen"
_ids = iter(range(1000, 10**6))


def _run(text: str, deleted: bool = False, rpr=None):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _change(tag: str, run):
    el = OxmlElement(tag)
    el.set(qn("w:id"), str(next(_ids)))
    el.set(qn("w:author"), AUTHOR)
    el.set(qn("w:date"), datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    el.append(run)
    return el


def _find(paragraphs, needle: str):
    """(paragraph, text, start, length) of the first paragraph whose text holds `needle` (raw, then normalized)."""
    for p in paragraphs:
        i = p.text.find(needle)
        if i >= 0:
            return p, p.text, i, len(needle)
    n = normalize(needle)
    for p in paragraphs:
        t = normalize(p.text)
        i = t.find(n)
        if i >= 0:
            return p, t, i, len(n)
    return None


# A clause number at the start of added text: "3.4 ", "12 ", "4.1.2 ".
_CLAUSE = re.compile(r"^\d+(\.\d+)*\.?\s")


def _shared_words(a: str, b: str) -> tuple[int, int]:
    """Lengths of the start and end `a` and `b` share, cut at word boundaries, never overlapping."""
    pre = len(os.path.commonprefix([a, b]))
    if pre < min(len(a), len(b)):
        pre = a.rfind(" ", 0, pre) + 1
    suf = len(os.path.commonprefix([a[pre:][::-1], b[pre:][::-1]]))
    if suf < min(len(a), len(b)) - pre:
        tail = a[len(a) - suf:]
        cut = tail.find(" ")
        suf = len(tail) - cut if cut >= 0 else 0
    return pre, suf


def _comment(doc, parts: list, start, end, text: str | None):
    """Attach a comment over start..end, or drop the empty range."""
    if text:
        c = doc.comments.add_comment(text=text, author=AUTHOR, initials="L")
        for el in (start, end):
            el.set(qn("w:id"), str(c.comment_id))
        ref_run = OxmlElement("w:r")
        ref = OxmlElement("w:commentReference")
        ref.set(qn("w:id"), str(c.comment_id))
        ref_run.append(ref)
        parts.insert(parts.index(end) + 1, ref_run)
    else:
        parts.remove(start)
        parts.remove(end)


def _new_paragraph_after(doc, p, text: str, comment: str | None):
    """A tracked new paragraph after `p`, with `p`'s paragraph style: the run is w:ins and so is its paragraph mark."""
    new_p = OxmlElement("w:p")
    ppr = p._p.find(qn("w:pPr"))
    ppr = copy.deepcopy(ppr) if ppr is not None else OxmlElement("w:pPr")
    mark = ppr.find(qn("w:rPr"))
    if mark is None:
        mark = OxmlElement("w:rPr")
        ppr.append(mark)
    ins_mark = OxmlElement("w:ins")
    ins_mark.set(qn("w:id"), str(next(_ids)))
    ins_mark.set(qn("w:author"), AUTHOR)
    ins_mark.set(qn("w:date"), datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    mark.append(ins_mark)
    new_p.append(ppr)
    rpr = p.runs[0]._r.find(qn("w:rPr")) if p.runs else None
    start, end = OxmlElement("w:commentRangeStart"), OxmlElement("w:commentRangeEnd")
    parts = [start, _change("w:ins", _run(text, rpr=rpr)), end]
    _comment(doc, parts, start, end, comment)
    for el in parts:
        new_p.append(el)
    p._p.addnext(new_p)


def converted_copy(doc_id: str) -> Document:
    doc = Document()
    for para in load_doc(doc_id)["paras"]:
        doc.add_paragraph(para["text"])
    return doc


def redline(src: str | None, edits: list[dict], out: str, doc_id: str | None = None) -> dict:
    if src and src.lower().endswith(".docx"):
        doc = Document(src)
    elif doc_id:
        doc = converted_copy(doc_id)
    else:
        raise KitError("a .docx source or --doc is required")
    applied, missing = 0, []
    for e in edits:
        frm, to = e.get("from", ""), e.get("to", "")
        hit = _find(doc.paragraphs, frm) if frm else None
        if not hit:
            missing.append(e.get("a") or frm[:40])
            continue
        p, text, i, length = hit
        frm_text = text[i: i + length]
        tail = to[len(frm_text):] if to.startswith(frm_text) else None
        if tail is not None and _CLAUSE.match(tail.strip()):
            # A new clause after the anchor ("Missing" fix): its own paragraph, inserted; the anchor stays as it is.
            _new_paragraph_after(doc, p, tail.strip(), e.get("comment"))
            applied += 1
            continue
        # Only the words that change are deleted and inserted: the shared start and end stay plain text.
        pre, suf = _shared_words(frm_text, to)
        old_mid, new_mid = frm_text[pre: len(frm_text) - suf], to[pre: len(to) - suf]
        before, after = text[: i + pre], text[i + len(frm_text) - suf:]
        # ponytail: the paragraph is rebuilt as plain runs with the first run's formatting, so a second
        # edit in the same paragraph replaces the first. One edit per paragraph is the playbook norm.
        rpr = p.runs[0]._r.find(qn("w:rPr")) if p.runs else None
        for r in list(p._p.iter(qn("w:r"))):
            r.getparent().remove(r)
        start, end = OxmlElement("w:commentRangeStart"), OxmlElement("w:commentRangeEnd")
        parts = [_run(before, rpr=rpr), start]
        if old_mid:
            parts.append(_change("w:del", _run(old_mid, True, rpr)))
        if new_mid:
            parts.append(_change("w:ins", _run(new_mid, rpr=rpr)))
        parts += [end, _run(after, rpr=rpr)]
        _comment(doc, parts, start, end, e.get("comment"))
        for el in parts:
            p._p.append(el)
        applied += 1
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)
    return {"ok": not missing, "applied": applied, "missing": missing, **verify(out)}


def verify(path: str) -> dict:
    doc = Document(path)
    body = doc.element.body
    return {
        "insertions": len(body.findall(".//" + qn("w:ins"))),
        "deletions": len(body.findall(".//" + qn("w:del"))),
        "comments": len(list(doc.comments)),
    }
