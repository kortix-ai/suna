"""python3 -m lumen_kit.selftest — stdlib unittest over the fixtures (run from lumen-kit/)."""
import json
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIX = os.path.join(HERE, "fixtures")
os.environ["LUMEN_DRIVE_ROOT"] = FIX
MATTER = os.path.join(FIX, "Matters", "m1")
VECTORS = os.path.join(HERE, "..", "shared", "cite-vectors.json")

from lumen_kit import core, search, tools, validate  # noqa: E402


def cli(*args, stdin=None):
    r = subprocess.run([sys.executable, os.path.join(HERE, "bin", "lumen"), *args], input=stdin, capture_output=True, text=True,
                       env={**os.environ, "LUMEN_DRIVE_ROOT": FIX})
    return r.returncode, json.loads(r.stdout)


class Anchors(unittest.TestCase):
    def test_parse(self):
        d = core.load_doc("docA")
        self.assertEqual(d["meta"]["file"], "msa-fixture.pdf")
        self.assertEqual([p["a"] for p in d["paras"]], ["p1.1", "p1.2", "p2.3"])
        self.assertEqual(d["paras"][2]["heading"], "11. Limitation of Liability")

    def test_read_anchor_context(self):
        r = core.read("docA", anchor="p1.2", context=1)
        self.assertEqual([p["a"] for p in r["paras"]], ["p1.1", "p1.2", "p2.3"])

    def test_bad_doc_id(self):
        with self.assertRaises(core.KitError):
            core.text_path("../etc/passwd")


class Cite(unittest.TestCase):
    def test_vectors(self):
        with open(VECTORS, encoding="utf-8") as f:
            for case in json.load(f)["cases"]:
                got = core.normalize(case["quote"]) in core.normalize(case["para"])
                self.assertEqual(got, case["ok"], case["name"])

    def test_codes(self):
        ok = {"doc_id": "docA", "a": "p2.3", "quote": "Neither party's liability under this Agreement shall be limited"}
        self.assertEqual(core.check_citation(ok), {"ok": True})
        self.assertEqual(core.check_citation({**ok, "a": "p1.1"}), {"ok": False, "code": "wrong_anchor", "fixed_anchor": "p2.3"})
        self.assertEqual(core.check_citation({**ok, "quote": "liability is unlimited"})["code"], "not_found")
        self.assertEqual(core.check_citation({**ok, "doc_id": "nope"})["code"], "unknown_doc")
        self.assertEqual(core.check_citation({**ok, "quote": "x" * 301})["code"], "too_long")


    def test_check_prints_the_chat_link(self):
        ok = {"doc_id": "docA", "a": "p2.3", "quote": "Neither party's liability under this Agreement shall be limited"}
        code, out = cli("cite", "check", stdin=json.dumps({"citations": [ok, {**ok, "quote": "liability is unlimited"}]}))
        self.assertEqual(out["results"][0]["link"],
                         "[1](/cite?d=docA&a=p2.3&q=Neither%20party%27s%20liability%20under%20this%20Agreement%20shall%20be%20limited)")
        self.assertNotIn("link", out["results"][1])


class Matter(unittest.TestCase):
    def test_session_context(self):
        os.environ["X_SESSION_CONTEXT"] = json.dumps({"lumen_matter_dir": MATTER})
        try:
            code, out = cli("docs")
            self.assertEqual((code, [d["doc_id"] for d in out["docs"]]), (0, ["docA", "docB"]))
            self.assertEqual(cli("matter"), (0, {"matter": MATTER}))
        finally:
            del os.environ["X_SESSION_CONTEXT"]
        self.assertEqual(cli("docs")[0], 2)


class Search(unittest.TestCase):
    def test_ranking_and_prefix(self):
        hits = search.search(MATTER, "indemnif* hold harmless")["hits"]
        self.assertEqual((hits[0]["doc_id"], hits[0]["a"]), ("docA", "p1.2"))

    def test_doc_filter(self):
        hits = search.search(MATTER, "governed laws", docs=["docB"])["hits"]
        self.assertEqual({h["doc_id"] for h in hits}, {"docB"})


class Validate(unittest.TestCase):
    row = {"schema": "review-row/1", "doc_id": "docB", "cells": {
        "term": {"value": "7 years", "note": "Clause on duration.", "confidence": "high",
                 "citations": [{"doc_id": "docB", "a": "p1.1", "quote": "continue for seven (7) years"}]},
        "liability_cap": {"value": None, "not_found": True, "note": "Checked all paragraphs.", "confidence": "high", "citations": []}}}

    def test_valid_row(self):
        self.assertEqual(tools.result_errors("review-row", self.row), [])

    def test_schema_errors(self):
        bad = json.loads(json.dumps(self.row))
        bad["cells"]["term"]["confidence"] = "certain"
        bad["cells"]["term"]["citations"][0]["a"] = "3.12"
        del bad["cells"]["term"]["note"]
        errs = validate.errors(bad, {"$ref": "review-row"})
        self.assertEqual(len(errs), 3, errs)

    def test_found_value_needs_citation_and_real_quote(self):
        bad = json.loads(json.dumps(self.row))
        bad["cells"]["term"]["citations"][0]["quote"] = "seven years"
        bad["cells"]["governing_law"] = {"value": "England", "note": "", "confidence": "high", "citations": []}
        errs = tools.result_errors("review-row", bad)
        self.assertTrue(any("not_found" in e for e in errs) and any("needs at least one citation" in e for e in errs), errs)

    def test_every_schema_loads(self):
        for name in ["citation", "review-row", "answer", "issues", "diligence-report", "draft", "done", "input"]:
            self.assertIn("$id", validate.schema(name))


class Results(unittest.TestCase):
    def test_write_done_cli(self):
        with tempfile.TemporaryDirectory() as batch:
            code, out = cli("result", "write", "--batch", batch, "--schema", "review-row", "--out", "rows/docB.json",
                            stdin=json.dumps(Validate.row))
            self.assertEqual((code, out["ok"]), (0, True))
            code, out = cli("result", "write", "--batch", batch, "--schema", "review-row", "--out", "../x.json",
                            stdin=json.dumps(Validate.row))
            self.assertEqual(code, 2)
            code, out = cli("result", "done", "--batch", batch, "--status", "succeeded")
            self.assertEqual((code, out["outputs"]), (0, ["out/rows/docB.json"]))
            with open(os.path.join(batch, "done.json")) as f:
                self.assertEqual(json.load(f)["status"], "succeeded")

    def test_done_refuses_invalid_output(self):
        with tempfile.TemporaryDirectory() as batch:
            os.makedirs(os.path.join(batch, "out"))
            with open(os.path.join(batch, "out", "bad.json"), "w") as f:
                json.dump({"schema": "review-row/1", "doc_id": "docB", "cells": {"x": {"value": 1}}}, f)
            code, out = cli("result", "done", "--batch", batch, "--status", "succeeded")
            self.assertEqual((code, out["ok"]), (3, False))
            self.assertFalse(os.path.exists(os.path.join(batch, "done.json")))


class Redline(unittest.TestCase):
    def test_round_trip(self):
        from docx import Document
        from lumen_kit import redline
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "nda.docx")
            d = Document()
            d.add_paragraph("4.2 The obligations shall continue for a period of seven (7) years from the date of each disclosure.")
            d.save(src)
            out = os.path.join(tmp, "redline.docx")
            r = redline.redline(src, [{"a": "p1.1", "from": "seven (7) years", "to": "three (3) years", "comment": "Rule NDA-01: max 3 years."}], out)
            self.assertEqual((r["ok"], r["insertions"], r["deletions"], r["comments"]), (True, 1, 1, 1))
            text = Document(out).paragraphs[0]._p.xml
            self.assertIn('w:author="Lumen"', text)
            self.assertIn("seven (7) years", text)

    def test_converted_copy_for_pdf(self):
        from lumen_kit import redline
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "r.docx")
            r = redline.redline(None, [{"a": "p2.3", "from": "shall be limited or capped", "to": "shall not exceed the fees paid", "comment": ""}], out, doc_id="docA")
            self.assertEqual((r["ok"], r["insertions"], r["deletions"]), (True, 1, 1))


class Draft(unittest.TestCase):
    def test_render_fills_fields_and_marks_the_rest(self):
        from unittest import mock
        from docx import Document
        from lumen_kit import tools
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "t.md")
            with open(tpl, "w") as f:
                f.write("# Mutual NDA\n\nMade on {{effective_date}} between {{party_a}} and {{party_b}}.\n\n## 1. Law\n1.1 Governed by {{governing_law}}.\n1.2 Courts of the same place.\n")
            out = os.path.join(tmp, "out", "draft.docx")
            # The image may lack pandoc: the python-docx renderer takes over.
            with mock.patch.object(tools.subprocess, "run", side_effect=FileNotFoundError("pandoc")):
                r = tools.draft_render(tpl, {"party_a": "Alpha Example Ltd", "party_b": "Beta Example Inc.", "effective_date": "2026-10-20"}, out)
            self.assertEqual(r["unfilled"], ["governing_law"])
            text = "\n".join(p.text for p in Document(out).paragraphs)
            self.assertIn("Alpha Example Ltd", text)
            self.assertIn("[Governing law \u2014 to confirm]", text)
            self.assertNotIn("[[CONFIRM", text)
            with open(r["markdown"]) as f:
                self.assertIn("[[CONFIRM: governing_law]]", f.read())  # the marker Lumen validates stays in the Markdown
            d = Document(out)
            self.assertEqual(d.paragraphs[0].style.name, "Heading 1")
            # One paragraph per numbered clause, never 1.1 and 1.2 in one <w:p>.
            clauses = [p.text for p in d.paragraphs if p.text[:3] in ("1.1", "1.2")]
            self.assertEqual(len(clauses), 2)
            marked = [r for p in d.paragraphs for r in p.runs if "to confirm" in r.text]
            self.assertEqual(marked[0].font.highlight_color, 7)  # WD_COLOR_INDEX.YELLOW
            h = d.styles["Heading 1"].font
            self.assertEqual((str(h.color.rgb), h.name), ("000000", "Times New Roman"))



class Report(unittest.TestCase):
    def test_report_docx_names_documents_and_pages(self):
        from unittest import mock
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            os.makedirs(os.path.join(tmp, "out"))
            with open(os.path.join(tmp, "input.json"), "w") as f:
                json.dump({"documents": [{"doc_id": "d1", "filename": "msa.pdf"}]}, f)
            with open(os.path.join(tmp, "out", "report.json"), "w") as f:
                json.dump({"schema": "diligence-report/1", "files": ["report.docx"],
                           "sections": [{"heading": "Liability", "body_md": "Uncapped [1](/cite?d=d1&a=p2.4&q=x).", "citations": []}],
                           "red_flags": [{"title": "Uncapped liability", "severity": "high", "doc_id": "d1", "citations": []}]}, f)
            with mock.patch.object(tools.subprocess, "run", side_effect=FileNotFoundError("pandoc")):
                r = tools.report_docx(tmp)
            text = "\n".join(p.text for p in Document(r["docx"]).paragraphs)
            self.assertIn("Uncapped (msa.pdf, p. 2).", text)
            self.assertIn("HIGH: Uncapped liability (msa.pdf)", text)
            self.assertNotIn("/cite?", open(r["markdown"]).read())


if __name__ == "__main__":
    unittest.main(verbosity=1)
